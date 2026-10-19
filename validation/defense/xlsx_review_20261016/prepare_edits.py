"""Read current workbooks; stage the requested edits without writing XLSX."""
import copy
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
WORK = Path(__file__).resolve().parent
schema = json.loads((WORK / 'before/defense_tables.json').read_text(encoding='utf-8'))
catalog = json.loads((WORK / 'before/defense_catalog.json').read_text(encoding='utf-8'))
books = {}
for file in schema['files']:
    wb = openpyxl.load_workbook(WORK / 'before' / file['name'], read_only=False, data_only=False)
    book = {'name': file['name'], 'sheets': []}
    for spec in file['sheets']:
        sheet = wb[spec['name']]
        headers = [cell.value for cell in sheet[1]]
        columns = copy.deepcopy(spec['columns'])
        order = [headers.index(col['label']) for col in columns]
        rows = [[row[i] for i in order] for row in list(sheet.values)[1:] if any(v is not None for v in row)]
        for col, i in zip(columns, order):
            col['width'] = sheet.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width
        table = list(sheet.tables.values())[0]
        book['sheets'].append({'name': spec['name'], 'columns': columns, 'rows': rows,
                               'table': {'name': table.name, 'style': table.tableStyleInfo.name}})
    wb.close()
    books[file['name'][:2]] = book

def note(book, key, value):
    next(s for s in book['sheets'] if s['name'] == '填写说明')['rows'].append([key, value])

def column(key, label, width=76):
    return {'key': key, 'label': label, 'type': 'text', 'required': True, 'width': width}

skills = books['02']['sheets'][0]
skills['columns'].insert(2, column('description', '完整描述'))
full = {
    '755': '释放圣盾，嘲讽周围目标并使自身受到伤害-50%，持续4秒。基础冷却10秒，嘲讽范围400。',
    '17': '蓄力1.6秒后释放1个火球，命中后对目标周围半径150范围造成100%攻击力的火系爆炸伤害。基础冷却6秒，施放距离1200。',
    '20301': '为施放距离900内生命比例最低的1名存活友方英雄施加荆棘护甲：提供目标最大生命5%的护盾，并恢复德鲁伊攻击力60%的生命。护甲存在期间，攻击该目标的怪物受到德鲁伊攻击力50%的自然反伤；基础护盾耗尽时护甲与反伤结束。持续10秒，基础冷却15秒。闪避、持续伤害与反伤不触发反伤。',
    '20002': '向目标方向射出3根箭矢，每根造成50%攻击力的冰系伤害。基础冷却4秒，施放距离900，基础散射角7度。',
    '1435': '立即传送至1名怪物位置并攻击目标，对其造成300%攻击力的电系伤害。基础冷却8秒，施放距离400。',
}
linked = {}
for option in catalog['options'].values():
    for sid in option.get('unlock', []):
        sid = str(abs(int(sid)))
        if option.get('description'):
            text = f"{option['title']}：{option['description']}"
            if text not in linked.setdefault(sid, []):
                linked[sid].append(text)
for row in skills['rows']:
    sid = row[0]
    raw = catalog['skills'][sid]
    description = full.get(sid)
    if description is None:
        parts = ['原配置未提供该记录的独立技能原文。']
        if linked.get(sid):
            parts.append('关联词条原文（含强化条件，不代表该记录单独拥有全部效果）：\n'+'\n'.join(linked[sid]))
        scalar = [('type', '行为类型'), ('cd', '冷却毫秒'), ('cast_range', '施放距离'), ('attack_range', '作用距离'),
                  ('damage', '伤害参数'), ('damage_range', '伤害范围'), ('statusChange', '施加状态与时长'), ('skillGroupType', '后续技能链')]
        parts.append('本记录参数摘要：'+'；'.join(f'{label}={json.dumps(raw[key], ensure_ascii=False)}' for key, label in scalar if raw.get(key) not in [None, [], '']))
        description = '\n'.join(parts)
    row.insert(2, description)
note(books['02'], '完整描述', '当前五名英雄初始技能使用现行规则的完整说明。其他历史/派生记录保留关联词条原文与本记录参数摘要，缺少独立原文时明确标注。此列导出为技能description；说明不会自动反写数值或替换词条表的说明，修改数值后请同步维护文案。')

monsters = books['05']['sheets'][0]
old_columns = monsters['columns']
keep = [i for i, col in enumerate(old_columns) if col['key'] not in {'label', 'name', 'desTra'}]
new_columns = [old_columns[0], column('title', '怪物名称', 26), column('description', '怪物描述')]+[old_columns[i] for i in keep if i != 0]
new_rows = []
for row in monsters['rows']:
    sid = row[0]
    monster = catalog['monsters'][sid]
    name = row[1] or monster['title']
    if name in ['', '0', None]:
        name = f'原配置未命名（ID{sid}）'
    description = monster['description']
    if description in ['', '0', None]:
        description = '原配置未提供怪物说明；当前1～50关未引用该记录。'
    new_rows.append([sid, name, description]+[row[i] for i in keep if i != 0])
monsters['columns'], monsters['rows'] = new_columns, new_rows
note(books['05'], '名称与描述', '“怪物名称”“怪物描述”直接导出到怪物title/description，替代旧的名称/说明文本编号。已有原文保持原样；历史空记录明确标注未命名/缺少说明。名称不代替怪物ID。')

levels = books['06']['sheets'][0]
for col in levels['columns']:
    if col['key'] == 'hp_add': col.update(label='生命成长额外等级', width=26)
    if col['key'] == 'atk_add': col.update(label='攻击成长额外等级', width=26)
note(books['06'], '生命成长额外等级（原修正加级）', '这是额外成长等级，不是百分比。普通怪基础生命结算：基础生命＋每级生命增量×（关卡ID＋本列）；之后继续计算刷怪表的累计生命成长、现有每3波成长及两层生命系数。已合并原章节修正，无需再加一次章节值。')
note(books['06'], '攻击成长额外等级（原修正加级）', '基础攻击＋每级攻击增量×（关卡ID＋本列＋截至当前波的累计攻击成长加级）；之后乘关卡攻击系数与该刷怪行局内攻击系数。周期产生的精英还会额外加入刷怪表指定的精英成长等级。')
note(books['06'], '加级示例', '例如：基础生命100、每级生命增量20、关卡ID为5、本列填3，则生命成长这一段为100＋20×(5＋3)＝260。填3表示额外按3级成长，既不是生命＋3，也不是×3。')

spawn = books['07']['sheets'][0]
rename = {'elite_every': ('普通怪精英替换周期（只）', 32), 'elite_growth': ('周期精英额外成长等级', 30), 'angle': ('固定精英/首领入场角度（度）', 36)}
for col in spawn['columns']:
    if col['key'] in rename:
        col['label'], col['width'] = rename[col['key']]
pools = copy.deepcopy(books['08']['sheets'][0])
pools['add'] = True
pools['table']['name'] = 'SpawnMonsterPools'
books['07']['sheets'].append(pools)
note(books['07'], '怪池成员子表', '本文件的“怪池成员”子表维护怪物组ID、成员顺序、怪物ID与抽取权重；“刷怪批次”的怪物组ID引用这里。原独立08文件已合并并删除。')
note(books['07'], '普通怪精英替换周期（只）', '仅普通波生效。本关普通刷怪累计每第N只替换成精英，不增加怪物总数，跨波累计，不含固定精英/首领。填40表示第40、80、120…只为精英；0表示关闭周期精英。跨波改变N时，以当前行N检查本关累计只数。')
note(books['07'], '周期精英额外成长等级', '仅上列产生的精英生效。填11表示计算该精英生命和攻击时，各自的成长等级再加11；不是×11或＋11%。第8/14波固定精英和第20波首领不读取本列。')
note(books['07'], '固定精英/首领入场角度（度）', '仅“固定精英”“最终首领”行生效。0＝下方，90＝左方，180＝上方，270＝右方。普通波沿用现有上方随机横坐标入场，此列不生效，原180数值保留。')
for key, value in books['08']['sheets'][-1]['rows']:
    note(books['07'], '怪池说明：'+str(key), value)

changed = [books[key] for key in ['01','02','03','04','05','06','07']]
for book in changed:
    book['preview'] = 'I1:K5' if book['name'].startswith('07') else 'A1:F6'
    spec_file = next(f for f in schema['files'] if f['name'] == book['name'])
    for sheet in book['sheets']:
        existing = next((s for s in spec_file['sheets'] if s['name'] == sheet['name']), None)
        if existing is None:
            existing = copy.deepcopy(next(f for f in schema['files'] if f['name'].startswith('08'))['sheets'][0])
            spec_file['sheets'].append(existing)
        existing['columns'] = [{k:v for k,v in col.items() if k != 'width'} for col in sheet['columns']]
    # Only requested main tables and notes are authored; all other tabs are imported intact.
    notes_only = book['name'][:2] in {'01','03','04'}
    book['sheets'] = [s for s in book['sheets'] if (not notes_only and s is book['sheets'][0]) or s['name'] in {'填写说明','怪池成员'}]
    for sheet in book['sheets']:
        if sheet['name'] == '填写说明':
            sheet['columns'][0]['width'] = 36
            sheet['columns'][1]['width'] = 92
            for row in sheet['rows']:
                row[1] = row[1].replace('保存八份表格', '保存修改后的七份表格')
schema['files'] = [f for f in schema['files'] if not f['name'].startswith('08')]
(WORK/'edits.json').write_text(json.dumps(changed, ensure_ascii=False, indent=2)+'\n', encoding='utf-8', newline='\n')
(WORK/'new_schema.json').write_text(json.dumps(schema, ensure_ascii=False, indent=2)+'\n', encoding='utf-8', newline='\n')
print('Prepared main-table edits and updated export notes; original numeric cells preserved.')
